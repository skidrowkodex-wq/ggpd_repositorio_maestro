# -*- coding: utf-8 -*-
import sys, os, re, unicodedata
sys.stdout.reconfigure(encoding='utf-8')
import openpyxl

PATH = r"C:\Users\innov\AppData\Local\Temp\opencode\tabla_unica_copia.xlsx"
wb = openpyxl.load_workbook(PATH, read_only=True, data_only=True)
print("HOJAS:", wb.sheetnames)
for name in wb.sheetnames:
    ws = wb[name]
    print(f"\n### {name!r}  max_row={ws.max_row} max_col={ws.max_column}")
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i >= 12:
            break
        vals = [("" if c is None else str(c))[:30] for c in row[:16]]
        print(f"  r{i+1:>3}: " + " | ".join(vals))
wb.close()

# formulas?
wb2 = openpyxl.load_workbook(PATH, read_only=True, data_only=False)
ws = wb2["TABLA UNICA"] if "TABLA UNICA" in wb2.sheetnames else wb2[wb2.sheetnames[0]]
print("\n--- buscando encabezados reales (primeras 30 filas) ---")
for i, row in enumerate(ws.iter_rows(values_only=True)):
    if i >= 30:
        break
    non = [c for c in row if c not in (None, "")]
    if len(non) >= 3:
        print(f"  r{i+1:>3}: " + " | ".join([("" if c is None else str(c))[:32] for c in row[:18]]))
wb2.close()
