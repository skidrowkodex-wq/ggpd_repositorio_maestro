# -*- coding: utf-8 -*-
"""
VERIFICACION: ¿el archivo "... - FAMILIA NORMALIZADO" contiene familias
actualizadas (mas recientes / mas completas) que core.cat_familias_materiales?
"""
import sys, os, re, json, unicodedata, urllib.request, datetime
from collections import Counter, defaultdict
sys.stdout.reconfigure(encoding='utf-8')
import openpyxl

ROOT = r"C:\Users\innov\OneDrive\Documents\Repositorio_Maestro"
ORI = os.path.join(ROOT, "TABLA_UNICA_MATERIALES_2026 - PRECIOS EURO.xlsx")
NRM = os.path.join(ROOT, "TABLA_UNICA_MATERIALES_2026 - PRECIOS EURO - FAMILIA NORMALIZADO.xlsx")


def norm(s):
    if s is None: return ""
    s = unicodedata.normalize("NFD", str(s))
    s = "".join(c for c in s if unicodedata.category(c) != "Mn").upper()
    s = re.sub(r"[^A-Z0-9]+", " ", s)
    return re.sub(r"\s+", " ", s).strip()


print("=" * 100)
print("1) FECHAS DE LOS ARCHIVOS")
print("=" * 100)
for p in (ORI, NRM):
    st = os.stat(p)
    print(f"  {os.path.basename(p)}")
    print(f"     {st.st_size:>9,} bytes | modificado: {datetime.datetime.fromtimestamp(st.st_mtime)}")

# ------------------------------------------------------------ BD
cfg = {}
for line in open(os.path.join(ROOT, ".env.local"), encoding="utf-8", errors="ignore"):
    m = re.match(r"\s*([A-Za-z0-9_]+)\s*=\s*(.*)$", line)
    if m: cfg[m.group(1)] = m.group(2).strip().strip('"').strip("'")


def get(p):
    req = urllib.request.Request(cfg["INSFORGE_URL"] + p)
    req.add_header("apikey", cfg["INSFORGE_API_KEY"])
    req.add_header("Authorization", "Bearer " + cfg["INSFORGE_API_KEY"])
    return json.load(urllib.request.urlopen(req, timeout=60))


fams = get("/api/database/records/cat_familias_materiales?schema=core"
           "&select=*&limit=200")
print("\n" + "=" * 100)
print("2) CATALOGO DE FAMILIAS EN LA BD (core.cat_familias_materiales)")
print("=" * 100)
print(f"  Total: {len(fams)} familias")
print(f"  Columnas: {list(fams[0].keys()) if fams else '-'}")
fams_bd = {f["codigo_familia"]: f for f in fams}
creados = [f.get("creado_en") for f in fams if f.get("creado_en")]
for f in sorted(fams, key=lambda x: x["codigo_familia"]):
    print(f"    {f['codigo_familia']:<40} {str(f.get('nombre_familia'))[:36]:<36}"
          f" creado_en={f.get('creado_en')}")
if creados:
    print(f"  Rango creado_en de las familias: {min(creados)}  ..  {max(creados)}")
else:
    print("  creado_en: sin valor (columna vacia)")

items = get("/api/database/records/cat_items_materiales_precios?schema=core"
            "&select=id,codigo_familia,creado_en&limit=2000")
print(f"\n  Items de precio en BD: {len(items)}")
ic = [i["creado_en"] for i in items if i.get("creado_en")]
if ic:
    print(f"  Rango creado_en de los items: {min(ic)}  ..  {max(ic)}")
cf = Counter(i["codigo_familia"] for i in items)
print(f"  Familias con items: {len(cf)} | items sin familia: {cf.get(None, 0)}")

# ------------------------------------------------------------ ARCHIVOS
def leer(path, hoja):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[hoja]
    hdr = [c.value for c in ws[1]]
    i_cod = hdr.index("familia_normalizada") if "familia_normalizada" in hdr else None
    i_des = hdr.index("familia_nor_descripcion") if "familia_nor_descripcion" in hdr else None
    i_f2 = hdr.index("FAMILIA2")
    cod, desc, f2 = Counter(), Counter(), Counter()
    if i_cod is not None:
        for r in ws.iter_rows(min_row=2, values_only=True):
            if all(v in (None, "") for v in r): continue
            if r[i_cod]: cod[r[i_cod]] += 1
            if i_des is not None and r[i_des]: desc[r[i_des]] += 1
            if r[i_f2]: f2[str(r[i_f2]).strip()] += 1
    wb.close()
    return hdr, cod, desc, f2


print("\n" + "=" * 100)
print("3) CONTENIDO DEL ARCHIVO NORMALIZADO (hoja TABLA UNICA)")
print("=" * 100)
if not os.path.exists(NRM):
    print("  !! El archivo NO existe")
    sys.exit()
hdr, cod, desc, f2 = leer(NRM, "TABLA UNICA")
print(f"  Encabezados: {hdr}")
print(f"\n  3a. Columna 'familia_normalizada' -> {len(cod)} valores distintos:")
for k, v in cod.most_common():
    print(f"       {v:>5}  {k}")
print(f"\n  3b. Columna 'familia_nor_descripcion' -> {len(desc)} valores distintos:")
for k, v in desc.most_common():
    print(f"       {v:>5}  {k}")
print(f"\n  3c. Columna 'FAMILIA2' (taxonomia original del libro) -> {len(f2)} valores distintos")

# comparar normalizado vs BD
cod_bd = set(fams_bd)
cod_archivo = set(cod)
print("\n" + "=" * 100)
print("4) COMPARACION: familias del archivo normalizado  vs  familias de la BD")
print("=" * 100)
print(f"  En archivo: {len(cod_archivo)}   En BD: {len(cod_bd)}")
nuevas = sorted(cod_archivo - cod_bd)
faltan = sorted(cod_bd - cod_archivo)
print(f"\n  En el ARCHIVO y NO en la BD  ({len(nuevas)}): {nuevas or 'NINGUNA'}")
print(f"  En la BD y NO en el ARCHIVO  ({len(faltan)}): {faltan or 'NINGUNA'}")

# descripcion coincidente?
desc_bd = {norm(f["nombre_familia"]): f["codigo_familia"] for f in fams}
dif_desc = []
for d, n in desc.items():
    if norm(d) not in desc_bd:
        dif_desc.append((d, n))
print(f"\n  Descripciones del archivo que NO coinciden con nombre_familia de la BD "
      f"({len(dif_desc)}): {dif_desc or 'NINGUNA'}")

# ------------------------------------------------------------ FAMILIA2 (fuente real de novedad)
print("\n" + "=" * 100)
print("5) LA 'NOVEDAD' REAL: taxonomia FAMILIA2 del libro vs BD")
print("=" * 100)
hoja_ori_hdr, _, _, f2_ori = leer(ORI, "TABLA UNICA")
bd_norm = {norm(f["codigo_familia"]) for f in fams} | {norm(f["nombre_familia"]) for f in fams}
en_bd, fuera = [], []
for v, n in f2.most_common():
    nn = norm(v)
    ok = nn in bd_norm or any(b in nn or nn in b for b in bd_norm if b)
    (en_bd if ok else fuera).append((v, n))
print(f"  FAMILIA2 distintos: {len(f2)}")
print(f"  con equivalente literal en BD: {len(en_bd)}")
print(f"  SIN equivalente literal en BD: {len(fuera)}  -> {sum(n for _, n in fuera)} filas")
for v, n in fuera:
    print(f"      {n:>5}  {v!r}")

print("\n" + "=" * 100)
print("6) CONCLUSION")
print("=" * 100)
print(f"  Familias nuevas en el archivo: {len(nuevas)}")
print(f"  Familias de la BD ausentes en el archivo: {len(faltan)}")
