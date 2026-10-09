# -*- coding: utf-8 -*-
import sys, re, unicodedata, json, os
from collections import Counter, OrderedDict
sys.stdout.reconfigure(encoding='utf-8')
import openpyxl

ROOT = r"C:\Users\innov\OneDrive\Documents\Repositorio_Maestro"
SRC = os.path.join(ROOT, "TABLA_UNICA_MATERIALES_2026 - PRECIOS EURO.xlsx")

def norm(s):
    if s is None: return ""
    s = unicodedata.normalize("NFD", str(s))
    s = "".join(c for c in s if unicodedata.category(c) != "Mn").upper()
    s = re.sub(r"[^A-Z0-9]+", " ", s)
    return re.sub(r"\s+", " ", s).strip()

wb = openpyxl.load_workbook(SRC, data_only=True)
ws = wb["TABLA UNICA"]
hdr = [c.value for c in ws[1]]
print("COLUMNAS:", hdr)
print("FILAS:", ws.max_row, "COLS:", ws.max_column)

i_fam = hdr.index("FAMILIA2")
i_des = hdr.index("DESCRIPCION DEL MATERIAL")
cnt = Counter()
vacias = 0
ejemplos = {}
n = 0
for r in ws.iter_rows(min_row=2, values_only=True):
    if all(v in (None, "") for v in r):
        continue
    n += 1
    f = r[i_fam]
    if f in (None, ""):
        vacias += 1
        cnt["(VACIO)"] += 1
        continue
    cnt[str(f).strip()] += 1
    ejemplos.setdefault(str(f).strip(), r[i_des])

print(f"\nFilas de datos: {n} | FAMILIA2 vacias: {vacias}")
print(f"Valores distintos de FAMILIA2: {len(cnt)}\n")
for k, v in cnt.most_common():
    print(f"  {v:>5}  {k!r}")
wb.close()

# familias canonicas de la BD
cfg = {}
for line in open(os.path.join(ROOT, ".env.local"), encoding="utf-8", errors="ignore"):
    m = re.match(r"\s*([A-Za-z0-9_]+)\s*=\s*(.*)$", line)
    if m: cfg[m.group(1)] = m.group(2).strip().strip('"').strip("'")
import urllib.request
req = urllib.request.Request(cfg["INSFORGE_URL"] + "/api/database/records/cat_familias_materiales?schema=core&select=codigo_familia,nombre_familia&limit=100")
req.add_header("apikey", cfg["INSFORGE_API_KEY"]); req.add_header("Authorization", "Bearer " + cfg["INSFORGE_API_KEY"])
fams = json.load(urllib.request.urlopen(req, timeout=60))
print("\n21 FAMILIAS CANONICAS (BD core):")
for f in fams:
    print(f"  {f['codigo_familia']:<40} {f['nombre_familia']}")
