# -*- coding: utf-8 -*-
import sys, os, datetime
from collections import Counter
sys.stdout.reconfigure(encoding='utf-8')
import openpyxl

ROOT = r"C:\Users\innov\OneDrive\Documents\Repositorio_Maestro"
ORI = os.path.join(ROOT, "NAC_2026_GGPD_INSTRUMENTO_LISTADO_MATERIALES_PLANES_V03_PROYECCION_2024_2028.xlsx")
DST = os.path.join(ROOT, "NAC_2026_GGPD_INSTRUMENTO_LISTADO_MATERIALES_PLANES_V03_PROYECCION_2024_2028 - FAMILIA NORMALIZADO.xlsx")

for p in (ORI, DST):
    st = os.stat(p)
    print(f"{'ORIGEN' if p==ORI else 'NUEVO '} {st.st_size:>9,} bytes  {datetime.datetime.fromtimestamp(st.st_mtime)}  {os.path.basename(p)[:60]}")

wb = openpyxl.load_workbook(DST, data_only=True)
print("\nHOJAS:", wb.sheetnames)
ws = wb["LISTADO"]
hdr = [c.value for c in ws[1]]
print("\nLISTADO columnas 23-25:", hdr[22:])
i_fam, i_d, i_c = hdr.index("FAM"), hdr.index("DESC"), hdr.index("fam_mat_normalizado")

cnt, vac, n = Counter(), 0, 0
ej = []
for r in ws.iter_rows(min_row=2, values_only=True):
    if all(v in (None, "") for v in r): continue
    n += 1
    cod = r[i_c]
    if not cod: vac += 1
    else:
        cnt[cod] += 1
        if len(ej) < 8: ej.append((str(r[i_fam])[:34], str(r[i_d])[:34], cod, r[hdr.index('fam_mat_descripcion')]))
print(f"\nFilas de datos: {n} | con familia: {sum(cnt.values())} | en blanco: {vac}")
print("\nMUESTRA (FAM | DESC | fam_mat_normalizado | fam_mat_descripcion):")
for e in ej:
    print(f"  {e[0]:<34} | {e[1]:<34} | {e[2]:<34} | {e[3]}")
print(f"\nDISTRIBUCION ({len(cnt)} familias):")
for k, v in cnt.most_common():
    print(f"  {v:>4}  {k}")

print("\n--- FAMILIAS_PENDIENTES ---")
wp = wb["FAMILIAS_PENDIENTES"]
for row in wp.iter_rows(min_row=4, values_only=True):
    print("   ", [str(c)[:46] if c is not None else "" for c in row[:4]])

print("\n--- RESUMEN_FAM_NORMALIZADA (top 12) ---")
wr = wb["RESUMEN_FAM_NORMALIZADA"]
for i, row in enumerate(wr.iter_rows(values_only=True), 1):
    if i > 16: break
    v = [(f"{c:.3f}" if isinstance(c, float) else str(c))[:40] if c is not None else "" for c in row[:4]]
    if any(v): print("   ", v)
wb.close()

# integridad de las hojas originales
wbo = openpyxl.load_workbook(ORI, data_only=True)
wbn = openpyxl.load_workbook(DST, data_only=True)
print("\nINTEGRIDAD DE HOJAS ORIGINALES:")
for s in wbo.sheetnames:
    a, b = wbo[s].max_row, wbn[s].max_row
    print(f"   {s:<24} {a:>5} -> {b:>5} filas  {'OK' if a==b else '!! DIFIERE'}")
wbo.close(); wbn.close()
