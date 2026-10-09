# -*- coding: utf-8 -*-
import sys, os, re, unicodedata
from collections import Counter
sys.stdout.reconfigure(encoding='utf-8')
import openpyxl

ROOT = r"C:\Users\innov\OneDrive\Documents\Repositorio_Maestro"
SRC = os.path.join(ROOT, "TABLA_UNICA_MATERIALES_2026 - PRECIOS EURO.xlsx")
DST = os.path.join(ROOT, "TABLA_UNICA_MATERIALES_2026 - PRECIOS EURO - FAMILIA NORMALIZADO.xlsx")

def formulas(path):
    wb = openpyxl.load_workbook(path)
    n = {}
    for s in wb.sheetnames:
        c = 0
        for row in wb[s].iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    c += 1
        n[s] = c
    wb.close()
    return n

fs, fd = formulas(SRC), formulas(DST)
print("FORMULAS por hoja  (original -> nuevo):")
for s in fs:
    marca = "OK" if fs[s] == fd.get(s) else "!! CAMBIO"
    print(f"  {s:<24} {fs[s]:>5} -> {fd.get(s,'-'):>5}   {marca}")

wb = openpyxl.load_workbook(DST, data_only=True)
print("\nHOJAS:", wb.sheetnames)
ws = wb["TABLA UNICA"]
hdr = [c.value for c in ws[1]]
print("\nCABECERA TABLA UNICA:", hdr)
print("Filas:", ws.max_row, "Cols:", ws.max_column)

i_cod, i_des = hdr.index("familia_normalizada"), hdr.index("familia_nor_descripcion")
cnt = Counter(); vac = 0; inval = set(); ejemplos = []
for r in ws.iter_rows(min_row=2, values_only=True):
    if all(v in (None, "") for v in r): continue
    cod = r[i_cod]; des = r[i_des]
    if not cod:
        vac += 1; continue
    cnt[cod] += 1
    if not des: inval.add(cod)
    if len(ejemplos) < 6:
        ejemplos.append((r[0], r[1], r[2], cod, des))

print(f"\nFilas con familia_normalizada: {sum(cnt.values())} | vacias: {vac} | sin descripcion: {inval or 'ninguna'}")
print("\nMUESTRA:")
print(f"  {'#':>4} | {'FAMILIA2':<28} | {'DESCRIPCION':<34} | {'NORMALIZADO':<34} | descripcion familia")
for e in ejemplos:
    print(f"  {str(e[0]):>4} | {str(e[1])[:28]:<28} | {str(e[2])[:34]:<34} | {str(e[3]):<34} | {e[4]}")

print(f"\nDISTRIBUCION ({len(cnt)} familias):")
for k, v in cnt.most_common():
    print(f"  {v:>5}  {k}")
print(f"  {sum(cnt.values()):>5}  TOTAL")

print("\n--- HOJA RESUMEN POR FAMILIA ---")
wr = wb["RESUMEN POR FAMILIA"]
for i, row in enumerate(wr.iter_rows(values_only=True), 1):
    if i > 32: break
    vals = [("" if c is None else (f"{c:.4f}" if isinstance(c, float) else str(c)))[:52] for c in row[:7]]
    if any(vals):
        print(f"  r{i:>2}: " + " | ".join(vals))
wb.close()
