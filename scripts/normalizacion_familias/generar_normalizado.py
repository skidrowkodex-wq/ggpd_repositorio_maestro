# -*- coding: utf-8 -*-
"""
Genera: TABLA_UNICA_MATERIALES_2026 - PRECIOS EURO - FAMILIA NORMALIZADO.xlsx
  - Hoja 'TABLA UNICA': + columna N 'familia_normalizada' (codigo) y O 'familia_nor_descripcion'
  - Hoja 'RESUMEN POR FAMILIA': conteo de materiales, % y suma de demanda/precio
Reglas acordadas:
  1) Si el material existe en BD (804 con codigo_familia) -> manda la BD.
  2) Si no -> mapeo FAMILIA2 (52) -> 21 familias canonicas.
  3) Rotulos sin familia MDM -> MATERIALES_VARIOS.
"""
import sys, os, re, json, shutil, unicodedata, urllib.request
from collections import Counter, defaultdict, OrderedDict
from copy import copy

sys.stdout.reconfigure(encoding='utf-8')
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

ROOT = r"C:\Users\innov\OneDrive\Documents\Repositorio_Maestro"
SRC = os.path.join(ROOT, "TABLA_UNICA_MATERIALES_2026 - PRECIOS EURO.xlsx")
DST = os.path.join(ROOT, "TABLA_UNICA_MATERIALES_2026 - PRECIOS EURO - FAMILIA NORMALIZADO.xlsx")
PROP = json.load(open(os.path.join(os.path.dirname(os.path.abspath(__file__)), "prop_familia2.json"), encoding="utf-8"))


def norm(s):
    if s is None: return ""
    s = unicodedata.normalize("NFD", str(s))
    s = "".join(c for c in s if unicodedata.category(c) != "Mn").upper()
    s = re.sub(r"[^A-Z0-9]+", " ", s)
    return re.sub(r"\s+", " ", s).strip()


# ---------------------------------------------------------------- BD
cfg = {}
for line in open(os.path.join(ROOT, ".env.local"), encoding="utf-8", errors="ignore"):
    m = re.match(r"\s*([A-Za-z0-9_]+)\s*=\s*(.*)$", line)
    if m: cfg[m.group(1)] = m.group(2).strip().strip('"').strip("'")

def get(path):
    req = urllib.request.Request(cfg["INSFORGE_URL"] + path)
    req.add_header("apikey", cfg["INSFORGE_API_KEY"])
    req.add_header("Authorization", "Bearer " + cfg["INSFORGE_API_KEY"])
    return json.load(urllib.request.urlopen(req, timeout=60))

items = get("/api/database/records/cat_items_materiales_precios?schema=core"
            "&select=descripcion_item,codigo_familia&limit=2000")
fams = OrderedDict((f["codigo_familia"], f["nombre_familia"]) for f in
                   get("/api/database/records/cat_familias_materiales?schema=core&select=codigo_familia,nombre_familia&limit=100"))
bd = {}
for it in items:
    bd.setdefault(norm(it["descripcion_item"]), it["codigo_familia"])
print(f"BD: {len(items)} items | {len(fams)} familias | {len(bd)} descripciones indexadas")

# ---------------------------------------------------------------- LIBROS
wbv = openpyxl.load_workbook(SRC, data_only=True)   # valores
wb = openpyxl.load_workbook(SRC)                    # preserva formulas -> se guarda

ws = wb["TABLA UNICA"]
wsv = wbv["TABLA UNICA"]
hdr = [c.value for c in ws[1]]
i_num, i_fam, i_des = hdr.index("#"), hdr.index("FAMILIA2"), hdr.index("DESCRIPCION DEL MATERIAL")
i_precio = hdr.index("PRECIO (EUR)")
NCOL_COD, NCOL_DES = len(hdr) + 1, len(hdr) + 2
print(f"'TABLA UNICA': {ws.max_row} filas x {ws.max_column} cols -> nuevas cols {NCOL_COD}/{NCOL_DES}")

# demanda desde AGREGADOS DE NAC (une por '#')
demanda = {}
if "AGREGADOS DE NAC" in wbv.sheetnames:
    wa = wbv["AGREGADOS DE NAC"]
    ha = [c.value for c in wa[1]]
    ia_n, ia_d = ha.index("#"), ha.index("DEMANDA 2026 (NAC)")
    for row in wa.iter_rows(min_row=2, values_only=True):
        if row[ia_n] in (None, ""): continue
        try: demanda[int(row[ia_n])] = float(row[ia_d] or 0)
        except Exception: pass
    print(f"Demanda 2026 disponible para {len(demanda)} numeros de #")

# ---------------------------------------------------------------- ESCRITURA
ws.cell(row=1, column=NCOL_COD, value="familia_normalizada")
ws.cell(row=1, column=NCOL_DES, value="familia_nor_descripcion")
estilo_hdr = ws.cell(row=1, column=1)._style
ws.cell(row=1, column=NCOL_COD)._style = copy(estilo_hdr)
ws.cell(row=1, column=NCOL_DES)._style = copy(estilo_hdr)

ORIGEN = Counter()
por_fam = defaultdict(lambda: {"n": 0, "precio": 0.0, "demanda": 0.0, "n_demanda": 0})
total_filas = 0

for r in range(2, ws.max_row + 1):
    num = wsv.cell(row=r, column=i_num + 1).value
    f2 = wsv.cell(row=r, column=i_fam + 1).value
    des = wsv.cell(row=r, column=i_des + 1).value
    pre = wsv.cell(row=r, column=i_precio + 1).value
    if f2 in (None, "") and des in (None, ""):
        continue
    total_filas += 1
    ndes = norm(des)
    if ndes and ndes in bd:
        cod = bd[ndes]
        ORIGEN["A_BD_MDM"] += 1
    else:
        cod = PROP.get(norm(str(f2).strip()), None)
        if cod is None:
            cod = "MATERIALES_VARIOS"
            ORIGEN["C_SIN_FAMILIA_MDM"] += 1
        else:
            ORIGEN["B_FAMILIA2"] += 1
    if cod not in fams:
        ORIGEN["INVALIDO"] += 1
        cod = "MATERIALES_VARIOS"

    c_cod = ws.cell(row=r, column=NCOL_COD, value=cod)
    c_des = ws.cell(row=r, column=NCOL_DES, value=fams[cod])
    c_cod._style = copy(ws.cell(row=r, column=i_fam + 1)._style)
    c_des._style = copy(ws.cell(row=r, column=i_fam + 1)._style)

    agg = por_fam[cod]
    agg["n"] += 1
    try: agg["precio"] += float(pre or 0)
    except Exception: pass
    try:
        d = demanda.get(int(num)) if num not in (None, "") else None
        if d:
            agg["demanda"] += d; agg["n_demanda"] += 1
    except Exception:
        pass

print(f"Filas procesadas: {total_filas}")
print("Origen de la asignacion:", dict(ORIGEN))

# autofiltro -> extender a las columnas nuevas
if ws.auto_filter and ws.auto_filter.ref:
    m = re.match(r"([A-Z]+)(\d+):([A-Z]+)(\d+)", ws.auto_filter.ref)
    if m:
        ws.auto_filter.ref = f"{m.group(1)}{m.group(2)}:O{m.group(4)}"
        print("Autofiltro extendido a:", ws.auto_filter.ref)

# ---------------------------------------------------------------- RESUMEN
NOMBRE = "RESUMEN POR FAMILIA"
if NOMBRE in wb.sheetnames:
    del wb[NOMBRE]
wr = wb.create_sheet(NOMBRE)
titulo = ["FAMILIA NORMALIZADA", "DESCRIPCION DE LA FAMILIA", "CANTIDAD DE MATERIALES",
          "% PARTICIPACION", "SUMA PRECIO (EUR)", "SUMA DEMANDA 2026", "FILAS CON DEMANDA"]
wr.append(["CORPOELEC - GGPD | Normalizacion de familias - TABLA UNICA MATERIALES 2026 (PRECIOS EURO)"])
wr.append([f"Fuente: TABLA UNICA_MATERIALES_2026 - PRECIOS EURO.xlsx | Regla: BD MDM > FAMILIA2 > MATERIALES_VARIOS"])
wr.append([])
wr.append(titulo)

hdr_fill = PatternFill("solid", fgColor="1F4E79")
hdr_font = Font(bold=True, color="FFFFFF", size=10)
thin = Side(style="thin", color="BFBFBF")
for c in wr[4]:
    c.fill = hdr_fill; c.font = hdr_font
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border = Border(left=thin, right=thin, top=thin, bottom=thin)

orden = sorted(por_fam.items(), key=lambda kv: -kv[1]["n"])
total_n = sum(v["n"] for _, v in orden)
for cod, v in orden:
    wr.append([cod, fams[cod], v["n"], (v["n"] / total_n) if total_n else 0,
               round(v["precio"], 2), v["demanda"], v["n_demanda"]])
wr.append(["TOTAL", "", total_n, 1,
           round(sum(v["precio"] for _, v in orden), 2),
           sum(v["demanda"] for _, v in orden),
           sum(v["n_demanda"] for _, v in orden)])

for row in wr.iter_rows(min_row=5, max_row=wr.max_row, max_col=len(titulo)):
    for c in row:
        c.border = Border(left=thin, right=thin, top=thin, bottom=thin)
row_tot = wr.max_row
for c in wr[row_tot]:
    c.font = Font(bold=True)
    c.fill = PatternFill("solid", fgColor="DDEBF7")
for r in range(5, wr.max_row + 1):
    wr.cell(row=r, column=4).number_format = "0.0%"
    wr.cell(row=r, column=5).number_format = '#,##0.00'
    wr.cell(row=r, column=6).number_format = '#,##0'
wr.column_dimensions["A"].width = 38
wr.column_dimensions["B"].width = 38
wr.column_dimensions["C"].width = 22
wr.column_dimensions["D"].width = 15
wr.column_dimensions["E"].width = 18
wr.column_dimensions["F"].width = 20
wr.column_dimensions["G"].width = 18
wr.freeze_panes = "A5"
wr.auto_filter.ref = f"A4:G{wr.max_row - 1}"

# anotacion de limitaciones
wr.append([])
wr.append(["NOTAS"])
wr["A" + str(wr.max_row)].font = Font(bold=True)
wr.append(["1. 'TABLA UNICA' no tiene columna de cantidad; la 'SUMA DEMANDA 2026' proviene de la hoja 'AGREGADOS DE NAC' (DEMANDA 2026 (NAC)), unida por el numero '#'."])
wr.append(["2. Las 21 familias provienen de core.cat_familias_materiales (BD InsForge ggpd-data-maestra-0002)."])
wr.append(["3. Cuando el material existe en core.cat_items_materiales_precios se usa su codigo_familia (autoridad MDM); si no, se usa el mapeo de FAMILIA2."])
wr.append(["4. Rotulos de FAMILIA2 sin equivalente en el MDM se asignan a MATERIALES_VARIOS."])

wb.save(DST)
wbv.close(); wb.close()
print(f"\nGUARDADO: {DST}")
print(f"Tamano: {os.path.getsize(DST):,} bytes")
