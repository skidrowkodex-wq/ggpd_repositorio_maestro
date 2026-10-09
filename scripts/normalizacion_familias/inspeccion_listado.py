# -*- coding: utf-8 -*-
"""Inspeccion de la hoja LISTADO del instrumento NAC 2026 V03 PROYECCION."""
import sys, os, unicodedata, re
from collections import Counter
sys.stdout.reconfigure(encoding='utf-8')
import openpyxl

ROOT = r"C:\Users\innov\OneDrive\Documents\Repositorio_Maestro"
SRC = os.path.join(ROOT, "NAC_2026_GGPD_INSTRUMENTO_LISTADO_MATERIALES_PLANES_V03_PROYECCION_2024_2028.xlsx")

print("EXISTS:", os.path.exists(SRC), "|", f"{os.path.getsize(SRC):,} bytes")

wb = openpyxl.load_workbook(SRC, data_only=True)
print("HOJAS:", wb.sheetnames)
for s in wb.sheetnames:
    ws = wb[s]
    print(f"   {s:<32} {ws.max_row:>6} filas x {ws.max_column} cols")

ws = wb["LISTADO"]
print("\n--- LISTADO: primeras 6 filas (todas las columnas) ---")
for i, row in enumerate(ws.iter_rows(min_row=1, max_row=6, values_only=True), 1):
    print(f" r{i}:", [str(v)[:30] if v is not None else None for v in row])

# detectar fila de encabezado real
hdr_row = None
for i, row in enumerate(ws.iter_rows(min_row=1, max_row=10, values_only=True), 1):
    vals = [str(v).upper() if v else "" for v in row]
    if "FAM" in vals and "DESC" in vals:
        hdr_row = i
        break
print("\nFILA DE ENCABEZADO DETECTADA:", hdr_row)
hdr = [c.value for c in ws[hdr_row]]
print("COLUMNAS:", hdr)

i_fam = hdr.index("FAM")
i_desc = hdr.index("DESC")
print(f"\nIndices: FAM={i_fam}  DESC={i_desc}  max_col={ws.max_column}")

# ya existen las columnas pedidas?
exis = [h for h in hdr if h and str(h).lower() in ("fam_mat_normalizado", "fam_mat_descripcion")]
print("Columnas pedidas ya presentes:", exis or "NINGUNA")

cnt_fam, cnt_desc, vac_fam, vac_desc, n = Counter(), Counter(), 0, 0, 0
for row in ws.iter_rows(min_row=hdr_row + 1, values_only=True):
    if all(v in (None, "") for v in row):
        continue
    n += 1
    f = row[i_fam] if i_fam < len(row) else None
    d = row[i_desc] if i_desc < len(row) else None
    if f in (None, ""):
        vac_fam += 1
    else:
        cnt_fam[str(f).strip()] += 1
    if d in (None, ""):
        vac_desc += 1
    else:
        cnt_desc[str(d).strip()] += 1

print(f"\nFilas de datos: {n}")
print(f"FAM  -> distintos: {len(cnt_fam)} | vacios: {vac_fam}")
print(f"DESC -> distintos: {len(cnt_desc)} | vacios: {vac_desc}")

print("\n--- VALORES DISTINTOS DE 'FAM' (conteo) ---")
for k, v in cnt_fam.most_common():
    print(f"  {v:>5}  {k!r}")

print("\n--- PRIMEROS 30 VALORES DE 'DESC' ---")
for k, v in list(cnt_desc.most_common(30)):
    print(f"  {v:>5}  {k[:90]!r}")

# formula present?
wbf = openpyxl.load_workbook(SRC)
wsf = wbf["LISTADO"]
f = sum(1 for row in wbf["LISTADO"].iter_rows() for c in row
        if isinstance(c.value, str) and c.value.startswith("="))
print(f"\nFormulas en LISTADO: {f} | autofiltro: {wsf.auto_filter.ref} | freeze: {wsf.freeze_panes}")
import zipfile, re as _re
z = zipfile.ZipFile(SRC)
names = z.namelist()
print("Graficos:", len([n for n in names if _re.search(r"chart", n, _re.I)]),
      "| Tablas dinamicas:", len([n for n in names if _re.search(r"pivot", n, _re.I)]),
      "| Tablas excel:", len([n for n in names if n.startswith("xl/tables/")]),
      "| Imagenes:", len([n for n in names if n.startswith("xl/media/")]))
wb.close(); wbf.close()
