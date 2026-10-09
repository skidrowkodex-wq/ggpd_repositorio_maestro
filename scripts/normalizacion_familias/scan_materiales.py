# -*- coding: utf-8 -*-
"""Reconocimiento de estructura de los archivos en materiales_sept_2026."""
import os, sys, io, json, traceback

sys.stdout.reconfigure(encoding='utf-8')

BASE = r"C:\Users\innov\OneDrive\Documents\Repositorio_Maestro\materiales_sept_2026"

import openpyxl, xlrd


def peek_xlsx(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        rows = []
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= 6:
                break
            rows.append([("" if c is None else str(c))[:40] for c in row[:14]])
        out.append({"sheet": ws.title, "dims": f"{ws.max_row}x{ws.max_column}", "head": rows})
    wb.close()
    return out


def peek_xls(path):
    wb = xlrd.open_workbook(path, on_demand=True)
    out = []
    for name in wb.sheet_names():
        ws = wb.sheet_by_name(name)
        rows = []
        for r in range(min(6, ws.nrows)):
            rows.append([str(ws.cell_value(r, c))[:40] for c in range(min(14, ws.ncols))])
        out.append({"sheet": name, "dims": f"{ws.nrows}x{ws.ncols}", "head": rows})
        wb.unload_sheet(name)
    return out


results = {}
for root, dirs, files in os.walk(BASE):
    for f in sorted(files):
        p = os.path.join(root, f)
        rel = os.path.relpath(p, BASE)
        if f.startswith("~$"):
            continue
        ext = os.path.splitext(f)[1].lower()
        if ext not in (".xlsx", ".xls"):
            results[rel] = {"tipo": ext, "nota": "no legible como hoja de calculo"}
            continue
        try:
            results[rel] = peek_xlsx(p) if ext == ".xlsx" else peek_xls(p)
        except Exception as e:
            results[rel] = {"error": f"{type(e).__name__}: {e}"}

print(json.dumps(results, ensure_ascii=False, indent=1))
