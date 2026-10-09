# -*- coding: utf-8 -*-
"""Comprobacion extra: integridad tras posible re-guardado + mapeo de OTROS"""
import sys, os, datetime
from collections import Counter, defaultdict
sys.stdout.reconfigure(encoding='utf-8')
import openpyxl

ROOT = r"C:\Users\innov\OneDrive\Documents\Repositorio_Maestro"
NRM = os.path.join(ROOT, "TABLA_UNICA_MATERIALES_2026 - PRECIOS EURO - FAMILIA NORMALIZADO.xlsx")

print("AHORA           :", datetime.datetime.now())
st = os.stat(NRM)
print("ARCHIVO         :", datetime.datetime.fromtimestamp(st.st_mtime), f"{st.st_size:,} bytes")

wb = openpyxl.load_workbook(NRM, data_only=True)
print("HOJAS           :", wb.sheetnames)
ws = wb["TABLA UNICA"]
hdr = [c.value for c in ws[1]]
i_f2, i_cod = hdr.index("FAMILIA2"), hdr.index("familia_normalizada")
n = vac = 0
cruce = defaultdict(Counter)
for r in ws.iter_rows(min_row=2, values_only=True):
    if all(v in (None, "") for v in r): continue
    n += 1
    if not r[i_cod]: vac += 1
    cruce[str(r[i_f2]).strip()][r[i_cod]] += 1
print(f"\nFilas: {n} | familia_normalizada vacia: {vac} | cols: {ws.max_column}")

print("\nCruce FAMILIA2 -> familia_normalizada (selectivos):")
for f in ["OTROS", "HERRAMIENTAS", "LUMINARIAS Y BOMBILLOS", "MEDIDORES DE ENERGÍA",
          "TUBERÍAS Y ACCESORIOS", "EQUIPOS DE SEGURIDAD", "CINTAS AISLANTES",
          "CONDENSADORES", "EQUIPOS DE SUB ESTACIONES", "INTERRUPTORES Y SECCIONADORES"]:
    print(f"   {f:<32} -> {dict(cruce.get(f, {}))}")

print("\nHojas con conteo de filas:")
for s in wb.sheetnames:
    print(f"   {s:<24} {wb[s].max_row:>5} filas x {wb[s].max_column} cols")
if "RESUMEN POR FAMILIA" in wb.sheetnames:
    wr = wb["RESUMEN POR FAMILIA"]
    print("\nRESUMEN POR FAMILIA (primeras filas):")
    for i, row in enumerate(wr.iter_rows(values_only=True), 1):
        if i > 8: break
        print("   ", [str(c)[:38] if c is not None else "" for c in row[:7]])
wb.close()
