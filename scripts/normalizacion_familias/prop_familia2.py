# -*- coding: utf-8 -*-
"""
Propuesta FAMILIA2 (52) -> familia normalizada (21) para TABLA UNICA.
Cruza ademas por DESCRIPCION contra los 804 items ya normalizados en la BD
para detectar discrepancias de la propuesta.
"""
import sys, re, os, json, unicodedata, urllib.request
from collections import Counter, defaultdict
sys.stdout.reconfigure(encoding='utf-8')
import openpyxl

ROOT = r"C:\Users\innov\OneDrive\Documents\Repositorio_Maestro"
SRC = os.path.join(ROOT, "TABLA_UNICA_MATERIALES_2026 - PRECIOS EURO.xlsx")

def norm(s):
    if s is None: return ""
    s = unicodedata.normalize("NFD", str(s))
    s = "".join(c for c in s if unicodedata.category(c) != "Mn").upper()
    s = re.sub(r"[^A-Z0-9]+", " ", s)
    return re.sub(r"\s+", " ", s).strip()

# ----------------------------------------------------------------- propuesta
PROP = {
    # coinciden exactas con el canon
    "CONECTORES": "CONECTORES",
    "FUSIBLES": "FUSIBLES",
    "TRANSFORMADORES": "TRANSFORMADORES",
    "CABLES Y CONDUCTORES": "CABLES_Y_CONDUCTORES",
    "HERRAJES": "HERRAJES",
    "POSTES": "POSTES",
    "INTERRUPTORES": "INTERRUPTORES",
    "SECCIONADORES": "SECCIONADORES",
    "RECONECTADORES": "RECONECTADORES",
    "CORTACORRIENTES": "CORTACORRIENTES",
    "PARARRAYOS": "PARARRAYOS",
    "BANCO DE BATERIAS": "BANCO_DE_BATERIAS",
    "EQUIPOS DE SUB ESTACIONES": "EQUIPOS_DE_SUB_ESTACIONES",
    "MATERIALES VARIOS": "MATERIALES_VARIOS",
    "REGULADORES": "REGULADORES",
    "RECTIFICADORES": "RECTIFICADORES",
    "TEIPES CINTAS Y OTROS AISLANTES": "TEIPES_CINTAS_Y_OTROS_AISLANTES",
    "AISLADORES Y COPAS": "AISLADORES_Y_COPAS",
    # sinonimos directos
    "EQUIPOSDESUBESTACIONES": "EQUIPOS_DE_SUB_ESTACIONES",
    "EQUIPOS DE SUBESTACIONES": "EQUIPOS_DE_SUB_ESTACIONES",
    "AISLADORES": "AISLADORES_Y_COPAS",
    "ALUMBRADO PUBLICO": "MATERIALES_ALUMBRADO_PUBLICO",
    "LUMINARIAS Y BOMBILLOS": "MATERIALES_ALUMBRADO_PUBLICO",
    "CINTAS AISLANTES": "TEIPES_CINTAS_Y_OTROS_AISLANTES",
    "ESTRUCTURAS Y POSTES": "POSTES",
    "ESTRUCTURASDESOPORTE": "POSTES",
    "TORNILLERIA Y HERRAJES": "HERRAJES",
    "EQUIPO DE TIERRA": "HERRAJES",
    "SISTEMADEPUESTAATIERRA": "HERRAJES",
    "TRANSFORMADORES DISTRIBUCION": "TRANSFORMADORES",
    "TRANSFORMADORES POTENCIA S E": "TRANSFORMADORES",
    "TRANSFORMADORES DE MEDIDA": "TRANSFORMADORES_DE_MEDIDA_TC_Y_TP",
    "BATERIAS": "BANCO_DE_BATERIAS",
    "RECONECTADORES Y PROTECTORES": "RECONECTADORES",
    "ACEITEDIELECTRICO": "MATERIALES_VARIOS",
    "SERVICIOSAUXILIARES": "EQUIPOS_DE_SUB_ESTACIONES",
    "BANCODECONDESADORES": "EQUIPOS_DE_SUB_ESTACIONES",
    "CONDENSADORES": "EQUIPOS_DE_SUB_ESTACIONES",
    "TRAMO LINEA": "CABLES_Y_CONDUCTORES",
    # revision de ingenieria / diccionario MATERIALES 2026 -> canon 21
    "MEDIDORES DE ENERGIA": "TRANSFORMADORES_DE_MEDIDA_TC_Y_TP",
    "INTERRUPTORES Y SECCIONADORES": "INTERRUPTORES",      # ver cotejo por descripcion
    "EQUIPOS DE SEGURIDAD": "MATERIALES_VARIOS",
    "PINTURAS Y QUIMICOS": "MATERIALES_VARIOS",
    "CANDADOS Y LLAVES": "MATERIALES_VARIOS",
    "CAJAS Y GABINETES": "MATERIALES_VARIOS",
    "SELLOS Y PRECINTOS": "MATERIALES_VARIOS",
    "OTROS": "MATERIALES_VARIOS",
    "CONSTRUCCION": "MATERIALES_VARIOS",
    "HERRAMIENTAS": "MATERIALES_VARIOS",                    # <-- decision pendiente
    "REEMPLAZO DE PALILLOS": "MATERIALES_VARIOS",           # <-- decision pendiente
    "SIN FAMILIA NAC": "MATERIALES_VARIOS",                 # <-- decision pendiente
}

# ----------------------------------------------------------------- BD 804
cfg = {}
for line in open(os.path.join(ROOT, ".env.local"), encoding="utf-8", errors="ignore"):
    m = re.match(r"\s*([A-Za-z0-9_]+)\s*=\s*(.*)$", line)
    if m: cfg[m.group(1)] = m.group(2).strip().strip('"').strip("'")

def get(path):
    req = urllib.request.Request(cfg["INSFORGE_URL"] + path)
    req.add_header("apikey", cfg["INSFORGE_API_KEY"])
    req.add_header("Authorization", "Bearer " + cfg["INSFORGE_API_KEY"])
    return json.load(urllib.request.urlopen(req, timeout=60))

items = get("/api/database/records/cat_items_materiales_precios?schema=core"
            "&select=descripcion_item,codigo_familia&limit=2000")
fams = {f["codigo_familia"]: f["nombre_familia"]
        for f in get("/api/database/records/cat_familias_materiales?schema=core&select=codigo_familia,nombre_familia&limit=100")}
bd = defaultdict(set)
for it in items:
    bd[norm(it["descripcion_item"])].add(it["codigo_familia"])
print(f"BD: {len(items)} items, {len(fams)} familias")

# ----------------------------------------------------------------- TABLA UNICA
wb = openpyxl.load_workbook(SRC, data_only=True)
ws = wb["TABLA UNICA"]
hdr = [c.value for c in ws[1]]
i_fam, i_des = hdr.index("FAMILIA2"), hdr.index("DESCRIPCION DEL MATERIAL")

cnt = Counter()
sin_prop = Counter()
coinc_bd = ok_bd = dis_bd = 0
dis_ej = []
total = 0
por_fam = Counter()
for r in ws.iter_rows(min_row=2, values_only=True):
    if all(v in (None, "") for v in r):
        continue
    total += 1
    f2 = str(r[i_fam]).strip()
    cnt[f2] += 1
    canon = PROP.get(norm(f2))
    if canon is None:
        sin_prop[f2] += 1
        continue
    por_fam[canon] += 1
    # cotejo por descripcion
    d = norm(r[i_des])
    if d in bd:
        if canon in bd[d]:
            ok_bd += 1
        else:
            dis_bd += 1
            if len(dis_ej) < 25:
                dis_ej.append((r[i_des], f2, canon, sorted(bd[d])))
        coinc_bd += 1

wb.close()

print(f"\nFilas de datos: {total}")
print(f"FAMILIA2 SIN PROPUESTA: {len(sin_prop)} valores")
for k, v in sin_prop.most_common():
    print(f"   {v:>4}  {k!r}")

print(f"\nCOTEJO POR DESCRIPCION contra los 804 de la BD:")
print(f"  Filas cuyo material existe en la BD : {coinc_bd}")
print(f"    - familia propuesta COINCIDE     : {ok_bd}")
print(f"    - familia propuesta DIFIERE      : {dis_bd}")
print("\n  EJEMPLOS DE DIFERENCIA (material | FAMILIA2 -> propuesta | familia BD):")
for e in dis_ej:
    print(f"    {str(e[0])[:42]:<42} | {e[1][:26]:<26} -> {e[2]:<34} | {','.join(e[3])}")

print(f"\nDISTRIBUCION PROPUESTA ({len(por_fam)} familias usadas de 21):")
for k, v in sorted(por_fam.items(), key=lambda x: -x[1]):
    print(f"  {v:>5}  {k:<38} {fams.get(k,'')}")
print(f"  {sum(por_fam.values()):>5}  TOTAL")
no_usadas = [f for f in fams if f not in por_fam]
print(f"\nFamilias canonicas SIN filas: {no_usadas}")

json.dump(PROP, open(os.path.join(os.path.dirname(os.path.abspath(__file__)), "prop_familia2.json"), "w"),
          ensure_ascii=False, indent=1)
