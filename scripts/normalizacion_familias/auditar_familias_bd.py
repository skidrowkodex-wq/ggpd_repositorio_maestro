# -*- coding: utf-8 -*-
"""
Auditoria: valores de familia en el Excel vs catalogo core.cat_familias_materiales (BD).
Revisa TODAS las columnas cuyo encabezado contiene 'FAMILIA' en las 8 hojas.
"""
import sys, os, re, json, unicodedata, urllib.request
from collections import Counter, defaultdict, OrderedDict
sys.stdout.reconfigure(encoding='utf-8')
import openpyxl

ROOT = r"C:\Users\innov\OneDrive\Documents\Repositorio_Maestro"
SRC = os.path.join(ROOT, "TABLA_UNICA_MATERIALES_2026 - PRECIOS EURO.xlsx")

def norm(s):
    if s is None: return ""
    s = unicodedata.normalize("NFD", str(s))
    s = "".join(c for c in s if unicodedata.category(c) != "Mn").upper()
    s = re.sub(r"[^A-Z0-9]+", " ", s)
    return re.sub(r"\s+", " ", s).strip()

# ------------------------------------------------------------------ BD
cfg = {}
for line in open(os.path.join(ROOT, ".env.local"), encoding="utf-8", errors="ignore"):
    m = re.match(r"\s*([A-Za-z0-9_]+)\s*=\s*(.*)$", line)
    if m: cfg[m.group(1)] = m.group(2).strip().strip('"').strip("'")

def get(p):
    req = urllib.request.Request(cfg["INSFORGE_URL"] + p)
    req.add_header("apikey", cfg["INSFORGE_API_KEY"])
    req.add_header("Authorization", "Bearer " + cfg["INSFORGE_API_KEY"])
    return json.load(urllib.request.urlopen(req, timeout=60))

fams_bd = get("/api/database/records/cat_familias_materiales?schema=core"
              "&select=codigo_familia,nombre_familia&limit=100")
nom_bd = {norm(f["nombre_familia"]): f for f in fams_bd}
cod_bd = {norm(f["codigo_familia"]): f for f in fams_bd}
print(f"BD core.cat_familias_materiales: {len(fams_bd)} familias\n")

# ------------------------------------------------------------------ EXCEL
wb = openpyxl.load_workbook(SRC, data_only=True)
columnas = []          # (hoja, idx_col, encabezado)
for sn in wb.sheetnames:
    ws = wb[sn]
    hdr = [c.value for c in ws[1]]
    for i, h in enumerate(hdr):
        if h and "FAMILIA" in str(h).upper():
            columnas.append((sn, i, str(h)))

print("COLUMNAS DE FAMILIA DETECTADAS:")
for sn, i, h in columnas:
    print(f"  - [{sn}] col {i+1}: {h}")

reco = defaultdict(Counter)        # (hoja,col) -> Counter(valores)
for sn, i, h in columnas:
    ws = wb[sn]
    for row in ws.iter_rows(min_row=2, values_only=True):
        if i >= len(row): continue
        v = row[i]
        if v in (None, ""): continue
        reco[(sn, i, h)][str(v).strip()] += 1
wb.close()

# ------------------------------------------------------------------ COMPARACION
def clasificar(valor):
    n = norm(valor)
    if n in nom_bd:
        return "EN_BD", nom_bd[n]["codigo_familia"]
    if n in cod_bd:
        return "EN_BD", cod_bd[n]["codigo_familia"]
    # tolerancia: el valor del excel esta contenido en un nombre de la BD (o viceversa)
    for k, f in nom_bd.items():
        if k and (k in n or n in k):
            return "EN_BD_TOLERANTE", f["codigo_familia"]
    return "NO_EN_BD", None

print("\n" + "=" * 96)
print("RESULTADO POR COLUMNA")
print("=" * 96)
global_no = Counter()       # valor -> apariciones totales (distinct)
global_det = {}
for key in reco:
    sn, i, h = key
    c = reco[key]
    en, tol, fuera = [], [], []
    for v, n in c.items():
        cat, cod = clasificar(v)
        if cat == "EN_BD": en.append((v, n, cod))
        elif cat == "EN_BD_TOLERANTE": tol.append((v, n, cod))
        else:
            fuera.append((v, n, None))
            global_no[v] += n
            global_det[v] = sn
    print(f"\n[{sn}] {h}")
    print(f"  valores distintos: {len(c)} | en BD: {len(en)} | tolerantes: {len(tol)} | NO EN BD: {len(fuera)}")
    if tol:
        for v, n, cod in sorted(tol, key=lambda x: -x[1]):
            print(f"     ~ {n:>5}  {v!r}  -> coincide con BD: {cod}")
    if fuera:
        for v, n, _ in sorted(fuera, key=lambda x: -x[1]):
            print(f"     ✗ {n:>5}  {v!r}")

print("\n" + "=" * 96)
print("VALORES DEL EXCEL AUSENTES EN EL CATALOGO DE LA BD (unicos, todas las hojas)")
print("=" * 96)
filas_afectadas = sum(global_no.values())
print(f"  Valores distintos ausentes: {len(global_no)}   Filas afectadas (suma de apariciones): {filas_afectadas}\n")
for v, n in global_no.most_common():
    print(f"  {n:>5}  {v!r}   (hoja: {global_det[v]})")

# ------------------------------------------------------------------ inverso
usados = set()
for key in reco:
    for v in reco[key]:
        cat, cod = clasificar(v)
        if cod: usados.add(cod)
ausentes_bd = [f for f in fams_bd if f["codigo_familia"] not in usados]
print("\n" + "=" * 96)
print("INVERSO: familias de la BD que NO aparecen en ningun valor del Excel")
print("=" * 96)
if ausentes_bd:
    for f in ausentes_bd:
        print(f"  ✗ {f['codigo_familia']:<40} {f['nombre_familia']}")
else:
    print("  (ninguna: las 21 familias de la BD al menos una vez en el libro)")
print(f"\nFamilias BD usadas: {len(usados)} / {len(fams_bd)}")
